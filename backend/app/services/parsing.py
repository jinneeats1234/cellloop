"""Turn uploaded partner files into text for the LLM, plus any machine-readable curves.

Supported: PDF, XLSX, CSV/TSV/TXT, and raw impedance exports (ZView .z, Gamry/BioLogic text
exports, or any delimited file with frequency / Z' / Z'' columns).
"""

import csv
import io
import re
from dataclasses import dataclass, field

ALLOWED_EXTENSIONS = {".pdf", ".xlsx", ".csv", ".tsv", ".txt", ".z", ".dta", ".json", ".md"}

# The browser-supplied Content-Type is never trusted: everything that isn't a real PDF or
# spreadsheet is stored and served as plain text, so an uploaded HTML/SVG file can't run as a page.
_BINARY_TYPES = {
    ".pdf": ("application/pdf", b"%PDF-"),
    ".xlsx": ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", b"PK\x03\x04"),
}
MAX_PDF_PAGES = 500
MAX_XLSX_UNCOMPRESSED = 200 * 1024 * 1024  # zip-bomb guard


def safe_content_type(filename: str) -> str:
    entry = _BINARY_TYPES.get(_ext(filename))
    return entry[0] if entry else "text/plain; charset=utf-8"


def check_signature(filename: str, data: bytes) -> str | None:
    """A file whose bytes don't match its extension is rejected (returns the reason)."""
    entry = _BINARY_TYPES.get(_ext(filename))
    if entry and not data.startswith(entry[1]):
        return f"This file isn't a valid {_ext(filename)} file (its contents don't match the extension)."
    return None


class UnreadableFileError(ValueError):
    """The upload can't be parsed (corrupt, encrypted or not what its extension claims)."""


@dataclass
class ParsedFile:
    text: str
    eis_data: dict | None = None
    iv_data: dict | None = None
    warnings: list[str] = field(default_factory=list)


def _ext(filename: str) -> str:
    m = re.search(r"(\.[A-Za-z0-9]+)$", filename)
    return m.group(1).lower() if m else ""


def parse_file(filename: str, data: bytes) -> ParsedFile:
    ext = _ext(filename)
    if ext == ".pdf":
        try:
            return _parse_pdf(data)
        except UnreadableFileError:
            raise
        except Exception as exc:  # pypdf raises many types for damaged/encrypted files
            raise UnreadableFileError("Couldn't read this PDF. It may be damaged or password-protected.") from exc
    if ext == ".xlsx":
        try:
            return _parse_xlsx(data)
        except UnreadableFileError:
            raise
        except Exception as exc:
            raise UnreadableFileError("Couldn't read this spreadsheet. Save it as .xlsx (not .xls or .xlsb) and re-upload.") from exc
    text = data.decode("utf-8", errors="replace")
    rows = _sniff_rows(text)
    eis, iv = _detect_curves(rows)
    return ParsedFile(text=text, eis_data=eis, iv_data=iv)


def _parse_pdf(data: bytes) -> ParsedFile:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        raise UnreadableFileError("This PDF is password-protected. Remove the password and re-upload.")
    pages = []
    for i, page in enumerate(reader.pages, start=1):
        if i > MAX_PDF_PAGES:
            break
        pages.append(f"[page {i}]\n{page.extract_text() or ''}")
    text = "\n\n".join(pages)
    warnings = [] if text.strip() else ["PDF has no extractable text (scanned?). OCR is required."]
    if len(reader.pages) > MAX_PDF_PAGES:
        warnings.append(f"Only the first {MAX_PDF_PAGES} of {len(reader.pages)} pages were read.")
    return ParsedFile(text=text, warnings=warnings)


def _parse_xlsx(data: bytes) -> ParsedFile:
    import zipfile

    from openpyxl import load_workbook

    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        if sum(i.file_size for i in zf.infolist()) > MAX_XLSX_UNCOMPRESSED:
            raise UnreadableFileError("This spreadsheet is too large once uncompressed.")
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts, eis, iv = [], None, None
    for ws in wb.worksheets:
        rows = [["" if c is None else str(c) for c in row] for row in ws.iter_rows(values_only=True)]
        rows = [r for r in rows if any(cell.strip() for cell in r)]
        parts.append(f"[sheet: {ws.title}]\n" + "\n".join("\t".join(r) for r in rows))
        e, v = _detect_curves(rows)
        eis, iv = eis or e, iv or v
    return ParsedFile(text="\n\n".join(parts), eis_data=eis, iv_data=iv)


def _sniff_rows(text: str) -> list[list[str]]:
    """Split delimited text into cells. Picks the delimiter per line (instrument exports often
    mix "key: value" preambles with a comma/tab data block). Only double quotes are quote
    characters, because Z' / Z'' column headers contain apostrophes."""
    rows = []
    for ln in text.splitlines():
        if not ln.strip():
            continue
        delim = next((d for d in ("\t", ",", ";") if d in ln), None)
        if delim:
            rows.append(next(csv.reader([ln], delimiter=delim, quotechar='"')))
        else:
            rows.append(re.split(r"\s{2,}|\s+(?=-?\d)", ln.strip()))
    return rows


def _to_float(s: str) -> float | None:
    try:
        return float(s.strip())
    except (ValueError, AttributeError):
        return None


_EIS_COLS = {
    "freq": re.compile(r"freq|^f\s*\(|hz", re.I),
    "zre": re.compile(r"z'(?!')|z_?re|zreal|re\s*\(z\)|real", re.I),
    "zim": re.compile(r"z''|-z''|z_?im|zimag|im\s*\(z\)|imag", re.I),
}
_IV_COLS = {
    "j": re.compile(r"current|^j\b|^i\b|a/cm", re.I),
    "v": re.compile(r"volt|^v\b|^u\b|potential|e\s*\(v\)", re.I),
}


def _find_header(rows: list[list[str]], patterns: dict[str, re.Pattern]) -> tuple[int, dict[str, int]] | None:
    for r_idx, row in enumerate(rows[:200]):
        found: dict[str, int] = {}
        for c_idx, cell in enumerate(row):
            for name, pat in patterns.items():
                if name not in found and pat.search(cell or ""):
                    found[name] = c_idx
                    break
        if len(found) == len(patterns):
            return r_idx, found
    return None


def _columns(rows: list[list[str]], start: int, cols: dict[str, int]) -> dict[str, list[float]]:
    out: dict[str, list[float]] = {k: [] for k in cols}
    for row in rows[start + 1:]:
        vals = {k: _to_float(row[i]) if i < len(row) else None for k, i in cols.items()}
        if any(v is None for v in vals.values()):
            if out[next(iter(cols))]:
                break  # end of numeric block
            continue
        for k, v in vals.items():
            out[k].append(v)  # type: ignore[arg-type]
    return out


def _detect_curves(rows: list[list[str]]) -> tuple[dict | None, dict | None]:
    eis = iv = None
    hdr = _find_header(rows, _EIS_COLS)
    if hdr:
        cols = _columns(rows, *hdr)
        if len(cols["freq"]) >= 5:
            zim = cols["zim"]
            header_cell = rows[hdr[0]][hdr[1]["zim"]]
            # Normalise to the physical sign convention (capacitive arcs have Z'' < 0).
            if header_cell.strip().startswith("-") or (zim and sum(zim) > 0):
                zim = [-z for z in zim]
            eis = {"freq_hz": cols["freq"], "z_real": cols["zre"], "z_imag": zim}
    hdr = _find_header(rows, _IV_COLS)
    if hdr:
        cols = _columns(rows, *hdr)
        if len(cols["j"]) >= 3:
            iv = {"current_density_a_cm2": cols["j"], "voltage_v": cols["v"]}
    return eis, iv
